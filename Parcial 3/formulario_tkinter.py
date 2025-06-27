import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os

nombre_archivo = "datos2.xlsx"
if os.path.exists(nombre_archivo):
    wb = load_workbook(nombre_archivo)
    ws = wb.active
else:
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre", "Carrera", "CURP", "Semestre", "No. Control", "Edad", "Sexo", "Dirección", "Teléfono", "Tipo de Sangre"])

def guardar_datos():
    datos = [entry.get() for entry in entradas]

    if any(not campo for campo in datos):
        messagebox.showwarning("Advertencia", "Todos los campos son obligatorios")
        return

    try:
        int(datos[5])  # Edad
        int(datos[8])  # Teléfono
    except ValueError:
        messagebox.showwarning("Advertencia", "Edad y Teléfono deben ser números")
        return

    ws.append(datos)
    wb.save(nombre_archivo)
    messagebox.showinfo("Información", "Los datos se han guardado con éxito")

    for entry in entradas:
        entry.delete(0, tk.END)

ventana = tk.Tk()
ventana.title("Registro de alumnos del CBTis 72")
ventana.configure(bg="blue")

label_style = {"bg": "blue", "fg": "white"}
entry_style = {"bg": "green", "fg": "black"}

campos = ["Nombre", "Carrera", "CURP", "Semestre", "No. Control",
          "Edad", "Sexo", "Dirección", "Teléfono", "Tipo de Sangre"]

entradas = []

# Mostrar 5 campos en la primera fila, 5 en la segunda
for i in range(5):
    label = tk.Label(ventana, text=campos[i], **label_style)
    label.grid(row=0, column=i, padx=5, pady=5)
    entry = tk.Entry(ventana, **entry_style)
    entry.grid(row=1, column=i, padx=5, pady=5)
    entradas.append(entry)

for i in range(5, 10):
    label = tk.Label(ventana, text=campos[i], **label_style)
    label.grid(row=2, column=i - 5, padx=5, pady=5)
    entry = tk.Entry(ventana, **entry_style)
    entry.grid(row=3, column=i - 5, padx=5, pady=5)
    entradas.append(entry)

boton_guardar = tk.Button(ventana, text="Guardar", command=guardar_datos, bg="white", fg="black")
boton_guardar.grid(row=4, column=0, columnspan=5, pady=10)

ventana.mainloop()

import subprocess
import platform

# Al final del script
if platform.system() == "Windows":
    subprocess.Popen(f'explorer /select,"{os.path.abspath(nombre_archivo)}"')
